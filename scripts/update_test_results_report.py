from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\Gigabyte\Downloads\Nhom13_BaoCao_TaiLieuThamKhao_ChinhSua.docx")
OUTPUT = Path(r"E:\IoT\Nhom13_BaoCao_TaiLieuThamKhao_ChinhSua_KetQuaKiemThu.docx")

source_doc = Document(SOURCE)
doc = Document(OUTPUT)
source_table = source_doc.tables[12]
table = doc.tables[12]

results = {
    "T01": "PASS (mô phỏng host): RTC hợp lệ lúc khởi động; callback NTP thành công; nguồn giờ chuyển RTC → NTP, giờ ICT UTC+7 đúng và RTC được cập nhật.",
    "T02": "PASS (mô phỏng host): RTC không hợp lệ, NTP không phản hồi; thử lại sau 10 giây, thời gian vẫn invalid và lịch mẫu không phát nhắc.",
    "T03": "PASS (mô phỏng host): Cấu hình và tồn kho đã lưu NVS; sau reboot giả lập, mất Wi-Fi, RTC/lịch A được khôi phục, tồn kho A = 42 và nhắc cục bộ lúc 08:00.",
    "T04": "PASS (mô phỏng host): Không có cache/NVS, RTC không hợp lệ và Wi-Fi mất; hệ thống không dùng lịch mẫu để phát nhắc.",
    "T05": "PASS (FSM + Firebase mock): TAKE, REMOVE đủ 2/2; servo tự đóng đúng mốc 8 giây; event VERIFIED_BY_SIMULATION; tồn kho A giảm 40 → 38 đúng một lần.",
    "T06": "PASS (FSM + Firebase mock): Không REMOVE tạo NO_DETECTION; REMOVE 1/2 hết 60 giây tạo MISMATCH_LESS; cả hai trường hợp không trừ tồn kho.",
    "T07": "PASS (FSM + Firebase mock): Lấy 2/1 tạo MISMATCH_MORE đúng một event; FSM báo lỗi và đóng servo; tồn kho B giữ nguyên 20.",
    "T08": "PASS (FSM + Firebase mock): Hoãn 5 phút rồi nhắc lại; không MISSED_DOSE trước đủ 10 phút sau nhắc lại; đến hạn chỉ tạo một event và userConfirmed = false.",
    "T09": "PASS (Firebase mock): Cấu hình liều A đổi 2 → 3 khi FSM rảnh; lần lấy kế tiếp xác minh 3/3, ghi event đúng liều và tồn kho 40 → 37.",
    "T10": "PASS (Firebase mock): Payload thiếu ngăn C và giờ 25:90 bị từ chối toàn bộ; cấu hình hợp lệ cũ của A vẫn là liều 2 lúc 08:00.",
    "T11": "PASS (Firebase mock): Lệnh đặt tồn kho 45 chuyển pending → applied, ghi NVS/Firebase và xóa lệnh; giá trị 101 bị rejected, tồn kho vẫn 45.",
    "T12": "PASS (Firebase mock): Hàng đợi tối đa 10 event và tồn kho 12 được lưu NVS khi offline; sau reboot giả lập, 10 event mới nhất được gửi và tồn kho cập nhật; không gửi trùng sau reboot tiếp.",
    "T13": "PASS (mô phỏng host): Hai ngăn cùng phút được xếp hàng; A phát trước, sau reboot B phát đúng một lần, không mất lượt hoặc phát lặp.",
}

expected_ids = set(results)
source_rows = {}
for row in source_table.rows[1:]:
    case_id = row.cells[0].text.strip()
    if case_id in expected_ids:
        source_rows[case_id] = [row.cells[i].text for i in range(3)]

case_ids = [f"T{i:02d}" for i in range(1, 14)]
if set(source_rows) != expected_ids:
    raise RuntimeError("Không lấy đủ mô tả T01–T13 từ báo cáo gốc.")

while len(table.rows) < len(case_ids) + 1:
    table.add_row()
while len(table.rows) > len(case_ids) + 1:
    table._tbl.remove(table.rows[-1]._tr)

headers = ["Mã", "Ca kiểm thử", "Kết quả mong đợi", "Kết quả thực tế"]
for col, value in enumerate(headers):
    table.cell(0, col).text = value

for row_idx, case_id in enumerate(case_ids, start=1):
    source_values = source_rows[case_id]
    values = [source_values[0], source_values[1], source_values[2], results[case_id]]
    for col_idx, value in enumerate(values):
        table.cell(row_idx, col_idx).text = value
table.cell(13, 2).text = "Hai ngăn cùng phút được xử lý qua reboot; không mất lượt hoặc phát lặp."


def shade_cell(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


widths = [Inches(0.55), Inches(1.38), Inches(2.85), Inches(1.75)]
table.autofit = False
for column, width in zip(table.columns, widths):
    column.width = width
for row_idx, row in enumerate(table.rows):
    prevent_split(row)
    for col_idx, cell in enumerate(row.cells):
        cell.width = widths[col_idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 or col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8)
                run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
                if row_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(31, 55, 79)
        if row_idx == 0:
            shade_cell(cell, "D9EAF7")
set_repeat_header(table.rows[0])


def replace_paragraph(index, text):
    paragraph = doc.paragraphs[index]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


replace_paragraph(
    313,
    "Phiên kiểm thử ngày 19/10/2026 đã chạy đủ 13 ca có thể mô phỏng (T01–T13); cả 13 ca đạt trên host simulation với các mô-đun firmware, NVS và Firebase REST mock. T14–T15 cần board hoặc nguồn thật nên được loại khỏi phạm vi mô phỏng theo yêu cầu. Tỷ lệ 13/13 PASS áp dụng cho bộ kiểm thử mô phỏng này.",
)
replace_paragraph(
    314,
    "Bộ test chạy trực tiếp mã schedule, time service, FSM và cloud_sync với RTC/NTP/Wi-Fi/Firebase giả lập có kiểm soát. Firebase thật đã nhận, trả lại và xóa một event thử có ID ngẫu nhiên; dữ liệu thử đã được dọn sạch. Wokwi không khởi động được qua CLI trong phiên này, nên chưa ghi nhận luồng Wokwi–Firebase–Dashboard đầu cuối.",
)
replace_paragraph(
    309,
    "Ảnh minh chứng bên dưới ghi lần build esp32dev có RAM 15,1% và flash 76,0%. Sau các cập nhật firmware để xử lý lỗi phát hiện trong kiểm thử, build cuối ngày 19/10/2026 vẫn SUCCESS (RAM 15,1%; flash 76,2%), được ghi tại mục 5.4. Các kết quả này chỉ xác nhận biên dịch firmware, chưa phải nghiệm thu Wokwi, Firebase/Dashboard hoặc board thật. [4]",
)
replace_paragraph(
    316,
    "Trong bảng, PASS là kết quả đạt tiêu chí trên host simulation; các tương tác Firebase trong bộ test dùng REST mock để kiểm tra payload và luồng đồng bộ. Smoke test Firebase thật được ghi riêng ở phần tổng kết. Các phép thử board thật T14–T15 không thuộc phạm vi mô phỏng.",
)
replace_paragraph(
    318,
    "Kết quả ngày 19/10/2026: 55 assertion FSM và 107 assertion tích hợp đều PASS; đủ T01–T13 đạt trong mô phỏng. PlatformIO build firmware esp32dev thành công, RAM 15,1% và flash 76,2%. Một event thử riêng đã PUT/GET/DELETE thành công trên Firebase thật; kiểm thử Wokwi–Dashboard đầu cuối chưa chạy.",
)
replace_paragraph(
    319,
    "Các kiểm thử đã xác nhận và xử lý các lỗi: NTP tiếp tục thử lại sau lỗi tạm thời; MISSED_DOSE không còn gắn userConfirmed sai; hàng đợi event ngoại tuyến được lưu NVS và chỉ xóa sau khi Firebase xác nhận; nhiều lịch cùng phút và SNOOZE vẫn đúng sau reboot hoặc tràn millis().",
)
replace_paragraph(
    324,
    "Các kết quả T01–T13 đã được chạy lại trên bộ mô phỏng host sau khi sửa firmware. Phần cần thiết bị thật như GPIO35, nguồn servo và thời lượng pin nằm ngoài phạm vi hiện tại; Dashboard đầu cuối cũng cần phiên Wokwi hoạt động để nghiệm thu cùng phiên firmware.",
)
replace_paragraph(
    329,
    "Chương kiểm thử ghi đủ 13 ca mô phỏng T01–T13 với kết quả 13/13 PASS trên host simulation. Hai ca T14–T15 cần board và nguồn thật nên được loại khỏi lần đánh giá này. Firebase thật đã qua kiểm tra ghi/đọc/xóa event riêng; tích hợp Wokwi–Dashboard đầu cuối chưa được xác nhận.",
)
replace_paragraph(
    157,
    "Điều kiện lịch gồm cấu hình hợp lệ, ngăn đã nạp, giờ hợp lệ và FSM sẵn sàng. Khi nhiều ngăn đến cùng phút, bộ lập lịch lưu hàng đợi và khóa từng ngăn theo ngày trong NVS; T13 xác nhận ngăn sau vẫn được nhắc sau reboot.",
)
replace_paragraph(
    172,
    "NVS lưu cache cấu hình, tồn kho, khóa lịch và hàng đợi sự kiện ngoại tuyến. Hàng đợi event có tối đa 10 phần tử; khi đầy, firmware bỏ event cũ nhất, còn khi Firebase xác nhận gửi thì event được xóa khỏi NVS. Vì vậy cần xét giới hạn hàng đợi và lỗi ghi NVS khi đánh giá khả năng lưu giữ dữ liệu.",
)
replace_paragraph(
    208,
    "Lịch so theo phút địa phương; trạng thái đã kích hoạt và hàng đợi các ngăn trùng phút được lưu trong NVS. Sau reboot trong phút hẹn, ngăn kế tiếp vẫn được lấy từ hàng đợi và không phát lại ngăn đã xử lý. T13 kiểm tra A/B cùng phút qua một lần khởi động lại; trường hợp thiết bị tắt đúng lúc đến lịch không được mô tả là có nhắc bù tự động.",
)
replace_paragraph(
    235,
    "Nhánh ngoại lệ phân loại MISSED_DOSE nếu không nhấn TAKE trong mười phút; NO_DETECTION nếu đã TAKE nhưng sau 60 giây không có REMOVE; và MISMATCH_LESS nếu số tín hiệu thấp hơn liều. SNOOZE tắt cảnh báo năm phút rồi khởi động lại mốc chờ mười phút. T08 xác nhận event MISSED_DOSE chỉ gửi một lần và userConfirmed=false.",
)
replace_paragraph(
    242,
    "Khả năng phục hồi gồm cache cấu hình/lịch, tồn kho và trạng thái chống lặp trong NVS; event chờ đồng bộ cũng được lưu NVS tối đa 10 phần tử. T12 thử mất mạng, reboot, kết nối lại và hàng đợi đầy; khi vượt giới hạn, firmware bỏ event cũ nhất. Nếu thiết bị tắt qua thời điểm hẹn thì không cam kết tự nhắc bù.",
)
replace_paragraph(
    243,
    "Khi mất Wi-Fi, nhắc và xác minh cục bộ tiếp tục nếu ESP32 còn nguồn, RTC hợp lệ và lịch đã được cache vào NVS. Event chưa gửi được lưu vào hàng đợi NVS tối đa 10 phần tử và khôi phục sau reboot; khi hàng đợi đầy, event cũ nhất bị loại. Pin DS3231 giữ giờ chứ không cấp điện cho cả thiết bị.",
)
replace_paragraph(
    262,
    "compartmentMatch=true phản ánh ngăn đang chạy theo cấu hình, không phải kết quả đo nhận dạng ngăn vật lý. Luồng MISSED_DOSE gọi verify_start(..., false), đặt cờ đã gửi event và không tạo thêm event trong nhánh ERROR; T08 xác nhận userConfirmed=false và chỉ có một bản ghi. Event gửi khi chưa có giờ hợp lệ vẫn cần được đánh giá riêng về timestamp.",
)
replace_paragraph(
    266,
    "Event gửi lỗi được giữ trong hàng đợi NVS tối đa 10 phần tử. Khi đầy, firmware bỏ event cũ nhất; khi gửi Firebase thành công, event đầu hàng đợi được xóa bền vững. T12 xác nhận hàng đợi được khôi phục sau reset và gửi lại khi online. Các thao tác REST vẫn dùng HTTP đồng bộ nên có thể chặn vòng lặp trong thời gian timeout; không mô tả toàn hệ thống là 100% non-blocking.",
)
replace_paragraph(
    272,
    "Khôi phục sau lỗi mạng dùng NVS cho lịch, tồn kho và tối đa 10 event chờ gửi. T12 xác nhận event và tồn kho qua reboot, đồng bộ khi có mạng, rồi không phát lại event đã ACK. Nếu vượt dung lượng, event cũ nhất bị loại; nếu tắt thiết bị qua thời điểm hẹn thì không có nhắc bù tự động.",
)
replace_paragraph(
    337,
    "Cache cấu hình trong NVS, RTC/NTP và FSM cho phép vận hành cục bộ khi mất mạng. Event chờ đồng bộ được lưu NVS tối đa 10 phần tử và được khôi phục sau reboot; các kết quả cùng giới hạn dung lượng được nêu rõ để đánh giá khách quan.",
)
replace_paragraph(
    343,
    "Các hướng nâng cấp tiếp theo gồm chuẩn hóa đơn vị timestamp lastSeen, giới hạn timeout HTTP, xác minh chứng chỉ TLS thay cho setInsecure(), và nghiệm thu Wokwi–Dashboard đầu cuối. Có thể bổ sung xử lý lỗi ghi NVS và chính sách hàng đợi khi vượt giới hạn sau khi đánh giá yêu cầu sử dụng.",
)
replace_paragraph(
    363,
    "□ Biên bản chạy lại T01–T13 trong mô phỏng kèm commit, log và PASS/FAIL; T14–T15 chỉ bổ sung nếu có board và nguồn thật",
)

doc.tables[14].cell(3, 1).text = (
    "Lịch ba ngăn A/B/C; các cữ cùng phút được xếp hàng và khóa chống lặp trong NVS; lịch phức tạp hơn chưa triển khai."
)
doc.tables[14].cell(4, 1).text = (
    "Lịch, tồn kho và tối đa 10 event chờ đồng bộ được lưu trong NVS qua reboot. Khi đầy, hàng đợi bỏ event cũ nhất; T12 đã kiểm tra tình huống này."
)
doc.tables[14].cell(4, 2).text = (
    "Event đã bị loại khi hàng đợi vượt giới hạn sẽ không thể khôi phục; cần theo dõi dung lượng và trạng thái đồng bộ."
)
doc.tables[15].cell(7, 2).text = (
    "Hàng đợi NVS tối đa 10 event, tồn tại qua reboot; vượt giới hạn sẽ bỏ event cũ nhất."
)
set_repeat_header(doc.tables[14].rows[0])
set_repeat_header(doc.tables[15].rows[0])
for row in doc.tables[14].rows:
    prevent_split(row)
for row in doc.tables[15].rows:
    prevent_split(row)

doc.paragraphs[317].paragraph_format.keep_with_next = True
doc.paragraphs[318].paragraph_format.keep_together = True
doc.paragraphs[325].paragraph_format.keep_with_next = True

summary = doc.tables[13]
summary_rows = [
    ["Nhóm kiểm thử", "Trạng thái tại 19/10/2026", "Ghi chú"],
    ["Logic firmware / sơ đồ", "162 assertion PASS", "Mã production; RTC, NTP, Wi-Fi, NVS và Firebase được điều khiển bằng mô phỏng host"],
    ["Kịch bản chức năng mô phỏng", "13/13 PASS (T01–T13)", "T14–T15 cần phần cứng thật nên không chạy theo phạm vi"],
    ["Build firmware esp32dev", "PASS", "PlatformIO; RAM 15,1%; flash 76,2%"],
    ["Firebase thật", "Event PUT/GET/DELETE PASS", "Event thử ID ngẫu nhiên đã được xóa; chưa nghiệm thu Wokwi–Dashboard đầu cuối"],
    ["Nguồn dự phòng 4 giờ", "Ngoài phạm vi mô phỏng", "T14–T15 cần board, nguồn và dữ liệu đo thực"],
]
for row_idx, values in enumerate(summary_rows):
    for col_idx, value in enumerate(values):
        summary.cell(row_idx, col_idx).text = value
summary.autofit = False
summary_widths = [Inches(1.65), Inches(2.25), Inches(2.2)]
for column, width in zip(summary.columns, summary_widths):
    column.width = width
for row_idx, row in enumerate(summary.rows):
    prevent_split(row)
    for col_idx, cell in enumerate(row.cells):
        cell.width = summary_widths[col_idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8)
                run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
                if row_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(31, 55, 79)
        if row_idx == 0:
            shade_cell(cell, "D9EAF7")
set_repeat_header(summary.rows[0])

doc.save(OUTPUT)
print(str(OUTPUT))
