from copy import copy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\Gigabyte\Downloads\nhóm13.docx")
OUTPUT = Path(r"E:\IoT\nhom13_cap_nhat_kiem_thu.docx")
shutil.copy2(SOURCE, OUTPUT)
doc = Document(OUTPUT)


def replace_paragraph(index, text):
    p = doc.paragraphs[index]
    if p.runs:
        p.runs[0].text = text
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(text)


replace_paragraph(
    299,
    "Ngày 19/10/2026, PlatformIO build cho esp32dev đã hoàn tất trên bản dự án và bản sao mô phỏng đã loại bỏ credentials thật. Đây là bằng chứng build PASS; kết quả build không thay cho kiểm thử chức năng, Firebase hay phần cứng thật.",
)
replace_paragraph(
    300,
    "Dự án sử dụng PlatformIO, board esp32dev, Arduino framework, Serial Monitor 115200 baud và diagram.json của Wokwi. Bản mô phỏng được chạy bằng include/credentials.example.h đã đổi tên thành credentials.h; không dùng thông tin Firebase thật. Kết quả phiên Wokwi ngày 19/10/2026 và các giới hạn tích hợp được ghi tại Chương 5.",
)
replace_paragraph(
    309,
    "Kết quả ngày 19/10/2026: pio run cho esp32dev thành công trên bản chính và bản sao test đã làm sạch credentials. Bản build chính dùng 49.328/327.680 byte RAM (15,1%) và 995.845/1.310.720 byte flash (76,0%). Bản sạch đã chạy trên Wokwi; log cho thấy RTC DS3231 không phản hồi lúc khởi động, sau đó NTP đồng bộ ESP32. Firebase trả HTTP 404 do dùng credentials mẫu nên chưa xác nhận đồng bộ Firebase–Dashboard; chưa kiểm thử trên board vật lý.",
)
replace_paragraph(
    313,
    "Bảng dưới đây giữ nguyên 15 kịch bản T01–T15 và thay hai cột “Kết quả cần quan sát”, “Môi trường” bằng kết quả thực tế và kết luận. Kết quả được ghi theo phiên build/Wokwi ngày 19/10/2026; chỉ T07 và phần xử lý cục bộ của T11 đạt. T05 không đạt trên đường kích hoạt mô phỏng; các ca còn lại được ghi chưa kiểm thử hoặc chưa nghiệm thu nếu thiếu điều kiện thật.",
)
replace_paragraph(
    314,
    "Wokwi chỉ chạy với bản credentials mẫu đã làm sạch. Firebase trả HTTP 404 ở /config/compartments và /inventory/compartments, vì vậy không thể xác nhận dữ liệu cloud hoặc Dashboard. Các kết luận về RTC, nguồn, GPIO và cơ khí trên board thật cần được chạy lại bằng phần cứng tương ứng.",
)
replace_paragraph(
    316,
    "Kết luận trong bảng dùng ba nhãn: ĐẠT cho phần đã quan sát đáp ứng kịch bản; KHÔNG ĐẠT cho phần đã chạy nhưng sai kỳ vọng; CHƯA KIỂM THỬ khi thiếu điều kiện hoặc bằng chứng. Trạng thái cloud được ghi riêng khi thao tác cục bộ đạt nhưng Firebase chưa được xác nhận.",
)
replace_paragraph(
    318,
    "Trong phiên 19/10/2026, firmware build PASS và bản mô phỏng sạch đã chạy trên Wokwi. Bằng chứng chức năng hiện có gồm ca quá liều T07 đạt ở FSM và LOAD T11 đạt ở xử lý tồn kho cục bộ; ca T05 không đạt trên lệnh kích hoạt mô phỏng. Các ca dựa vào Firebase/Dashboard, lịch lưu bền vững, board thật và nguồn dự phòng chưa được nghiệm thu.",
)
replace_paragraph(
    319,
    "Log Wokwi ghi Firebase HTTP 404 vì bản test dùng credentials mẫu; các event được đưa vào hàng đợi offline và chưa xác nhận gửi thành công. T01 ghi RTC DS3231 không phản hồi, dù NTP sau đó đồng bộ ESP32. Với T05, phím test 1 tạo IDLE → REMINDER → ERROR (MISSED_DOSE) ngay trước TAKE/REMOVE; cần rà soát mốc millis của helper mô phỏng và FSM trước khi kết luận lịch thật. Các phát hiện rà soát trước đây về lịch trùng giờ, cờ userConfirmed của MISSED_DOSE và đơn vị lastSeen vẫn cần được kiểm chứng riêng.",
)


test_results = [
    ("T01", "Wokwi: RTC DS3231 không phản hồi lúc khởi động; firmware báo TIME_INVALID rồi NTP đồng bộ ESP32. Chưa xác nhận giờ lấy từ RTC.", "CHƯA ĐẠT (RTC); NTP fallback đạt"),
    ("T02", "Lúc khởi động ban đầu có TIME_INVALID, nhưng NTP sau đó thành công; chưa giữ cả RTC và NTP ở trạng thái lỗi để chạy đủ ca.", "CHƯA KIỂM THỬ"),
    ("T03", "Chưa chạy mất Wi‑Fi và khởi động lại với cache lịch/RTC đã lưu.", "CHƯA KIỂM THỬ"),
    ("T04", "Log ghi chưa có lịch cloud/NVS hợp lệ và không dùng giờ mẫu; Wi‑Fi vẫn kết nối, Firebase trả 404 nên chưa tạo được điều kiện offline đúng yêu cầu.", "CHƯA KIỂM THỬ (offline)"),
    ("T05", "Wokwi, phím 1: IDLE → REMINDER → ERROR (MISSED_DOSE) ngay; chưa vào TAKE/REMOVE và chưa kiểm tra đóng nắp, event hoặc trừ kho.", "KHÔNG ĐẠT (đường kích hoạt mô phỏng)"),
    ("T06", "Không chạy đủ 60 giây với liều thiếu/không lấy; ca T05 dừng ở MISSED_DOSE trước bước xác minh liều.", "CHƯA KIỂM THỬ"),
    ("T07", "Lệnh O ghi nhận 2/1 viên, FSM IDLE → ERROR, LCD báo thừa liều và servo mở rồi đóng. Gửi event Firebase gặp HTTP 404, chuyển vào hàng đợi.", "ĐẠT (FSM Wokwi); cloud chưa xác nhận"),
    ("T08", "Không vào được bước SNOOZE vì trigger mô phỏng chuyển thẳng sang MISSED_DOSE; chưa chờ 5 phút hoặc kiểm tra event.", "CHƯA KIỂM THỬ"),
    ("T09", "Không có Firebase/Dashboard hợp lệ để đổi liều và kiểm tra firmware áp dụng cấu hình mới.", "CHƯA KIỂM THỬ"),
    ("T10", "Không gửi được cấu hình thiếu/sai lên Firebase; tính năng giữ cấu hình cũ chưa được chạy.", "CHƯA KIỂM THỬ"),
    ("T11", "Lệnh LOAD trên Wokwi tăng tồn kho ngăn A từ 10 lên 20 (+10); log ghi xếp lịch đồng bộ Firebase. Đồng bộ cloud chưa xác nhận do HTTP 404.", "ĐẠT (cục bộ); cloud chưa xác nhận"),
    ("T12", "HTTP 404 làm event vào hàng đợi offline; số event chờ tăng trong log. Đây không phải thử ngắt/kết nối lại có kiểm soát; chưa xác nhận phục hồi sau khi có mạng.", "CHƯA KIỂM THỬ (phục hồi)"),
    ("T13", "Chưa nạp hai lịch hợp lệ cùng giờ hoặc reboot trong phút hẹn; Firebase không trả cấu hình.", "CHƯA KIỂM THỬ"),
    ("T14", "Chưa đo GPIO35, nhiễu nút hoặc brownout servo trên board thật. Wokwi không thay thế phép đo điện.", "CHƯA KIỂM THỬ (board thật)"),
    ("T15", "Chưa có UPS/battery và chưa chạy thử bốn giờ; không có dữ liệu dòng, áp hoặc số lần reset.", "CHƯA KIỂM THỬ"),
]

table = doc.tables[12]
headers = ["Mã", "Ca kiểm thử", "Kết quả thực tế", "Kết luận"]
for i, header in enumerate(headers):
    table.cell(0, i).text = header
for row_idx, values in enumerate(test_results, start=1):
    for col_idx, value in enumerate(values):
        table.cell(row_idx, col_idx).text = value


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


widths = [Inches(0.46), Inches(1.78), Inches(2.91), Inches(0.95)]
table.autofit = False
for col, width in zip(table.columns, widths):
    col.width = width
for row_idx, row in enumerate(table.rows):
    prevent_row_split(row)
    for col_idx, cell in enumerate(row.cells):
        cell.width = widths[col_idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in (0, 3) or row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8)
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                if row_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(31, 55, 79)
        if row_idx == 0:
            set_cell_shading(cell, "D9EAF7")
set_repeat_table_header(table.rows[0])


summary = doc.tables[13]
summary_rows = [
    ["Nhóm kiểm thử", "Trạng thái bằng chứng tại 19/10/2026", "Ghi chú"],
    ["Logic firmware / sơ đồ", "Có trong mã và đã smoke test trên Wokwi", "Không thay thế nghiệm thu phần cứng thật"],
    ["Kịch bản chức năng Wokwi", "T07 và phần cục bộ T11 đạt; T05 không đạt", "T01/NTP và các ca còn lại xem bảng T01–T15"],
    ["Build firmware esp32dev", "PASS", "pio run trên bản chính và bản sạch; RAM 15,1%, flash 76,0%"],
    ["Wokwi → Firebase → Dashboard", "Chưa nghiệm thu đầu-cuối", "Bản sạch dùng credentials mẫu; endpoint Firebase trả HTTP 404"],
    ["Nguồn dự phòng 4 giờ", "Chưa nghiệm thu", "Chưa có UPS/battery và dữ liệu đo"],
]
for row_idx, values in enumerate(summary_rows):
    for col_idx, value in enumerate(values):
        summary.cell(row_idx, col_idx).text = value
summary.autofit = False
summary_widths = [Inches(1.65), Inches(2.25), Inches(2.2)]
for col, width in zip(summary.columns, summary_widths):
    col.width = width
for row_idx, row in enumerate(summary.rows):
    prevent_row_split(row)
    for col_idx, cell in enumerate(row.cells):
        cell.width = summary_widths[col_idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8)
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                if row_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(31, 55, 79)
        if row_idx == 0:
            set_cell_shading(cell, "D9EAF7")
set_repeat_table_header(summary.rows[0])

doc.save(OUTPUT)
print(OUTPUT)
