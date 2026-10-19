from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


SOURCE = Path(r"C:\Users\Gigabyte\Downloads\Bao_cao_IoT_Hop_thuoc_thong_minh_Nhom13_Ban_mo_rong.docx")
OUTPUT = Path(r"E:\IoT\Bao_cao_IoT_Hop_thuoc_thong_minh_Nhom13_Ban_mo_rong_huong_dan_hinh_anh.docx")


def add_labeled(doc, label, text):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(label + " ")
    run.bold = True
    paragraph.add_run(text)
    return paragraph


def add_item(doc, code, title, location, capture, caption, condition=None):
    heading = doc.add_paragraph(style="Heading 2")
    heading.paragraph_format.keep_with_next = True
    heading.add_run(f"{code}. {title}")
    add_labeled(doc, "Vị trí chèn:", location)
    add_labeled(doc, "Cách lấy ảnh:", capture)
    add_labeled(doc, "Chú thích gợi ý:", caption)
    if condition:
        add_labeled(doc, "Điều kiện sử dụng:", condition)


doc = Document(SOURCE)
page_break = doc.add_paragraph()
page_break.add_run().add_break(WD_BREAK.PAGE)
doc.add_heading("PHỤ LỤC – HƯỚNG DẪN BỔ SUNG HÌNH ẢNH", level=1)

doc.add_paragraph(
    "Các mục dưới đây đối chiếu trực tiếp với nội dung bản báo cáo mở rộng. "
    "Ảnh mô tả giao diện và sơ đồ có thể chèn ngay sau khi kiểm tra độ rõ. "
    "Ảnh kết quả kiểm thử chỉ được dùng khi nhóm đã chạy đúng ca thử trên phiên bản nộp và ghi nhận kết quả thực tế."
)

doc.add_heading("1. Ảnh đã có để chọn và chèn", level=2)
doc.add_paragraph(
    "Trong thư mục report_screenshots của dự án E:\\IoT hiện có bốn ảnh chụp: "
    "01-platformio-build-success.png (kết quả build), "
    "02-wokwi-mo-phong-mach.png (mô phỏng đang chạy), "
    "03-wokwi-so-do-linh-kien.png (sơ đồ linh kiện), "
    "04-wokwi-rtc-firebase-log.png (log NTP và tải cấu hình). "
    "Trước khi dùng làm minh chứng kiểm thử, đối chiếu ảnh với commit, thời điểm và cấu hình của bản sẽ nộp."
)

doc.add_heading("2. Ảnh nên bổ sung cho phần mô tả hệ thống", level=2)
add_item(
    doc, "H01", "Sơ đồ Wokwi và ánh xạ phần cứng",
    "Sau mục 2.2 hoặc đầu mục 3.2, cạnh bảng GPIO.",
    "Mở E:\\IoT trong VS Code, chọn diagram.json để hiện Wokwi Diagram Editor. "
    "Phóng to hoặc dùng phím F để vừa khung; bảo đảm ESP32, LCD, RTC, DHT22, servo, LED, buzzer và bảy nút đều nhìn rõ. "
    "Chụp vùng sơ đồ bằng công cụ chụp màn hình Windows. Có thể chọn ảnh 03 đã có.",
    "Hình: Sơ đồ mô phỏng phần cứng hộp thuốc thông minh trên Wokwi.",
    "Không dùng hình này làm bằng chứng có UPS, cảm biến đếm viên thật hoặc ba khóa ngăn độc lập."
)
add_item(
    doc, "H02", "Mạch Wokwi khi firmware đang chạy",
    "Trong mục 4.2, sau phần mô tả LCD, LED, servo và các nút.",
    "Biên dịch PlatformIO, mở diagram.json và bấm nút Play màu xanh. "
    "Chờ LCD hiện giờ/trạng thái; đặt khung sao cho mạch và các nút TAKE, SNOOZE, REMOVE, LOAD, MENU, UP, DOWN đọc được. "
    "Chụp cửa sổ Wokwi Simulator. Có thể chọn ảnh 02 đã có.",
    "Hình: Mô phỏng hộp thuốc thông minh đang vận hành trên Wokwi."
)
add_item(
    doc, "H03", "Kết quả biên dịch PlatformIO",
    "Sau mục 4.5 hoặc trong mục 5.2.",
    "Trong VS Code chạy PlatformIO Build cho môi trường esp32dev (nút dấu tích hoặc lệnh pio run). "
    "Chụp Terminal khi hiện rõ môi trường build, dòng SUCCESS hoặc lỗi và thời gian chạy. "
    "Có thể dùng ảnh 01 đã có nếu đó là đúng phiên bản nộp.",
    "Hình: Kết quả biên dịch firmware cho ESP32 bằng PlatformIO.",
    "Nếu dùng ảnh SUCCESS, cập nhật câu đánh giá trong mục 5.2 và ghi commit tương ứng; không giữ kết luận 'chưa xác nhận' cho chính lần build này."
)
add_item(
    doc, "H04", "Giao diện Dashboard tổng quan",
    "Trong mục 4.4.1, ngay sau đoạn mô tả các vùng giám sát.",
    "Mở dashboard/index.html của dự án trong trình duyệt; chờ Firebase cập nhật. "
    "Chụp phần tiêu đề, trạng thái kết nối, nhiệt độ/độ ẩm, trạng thái máy, ba thẻ ngăn, tồn kho và lịch sử sự kiện; "
    "nếu không vừa một khung, chụp hai ảnh theo phần trên và dưới trang.",
    "Hình: Dashboard giám sát trạng thái thiết bị và ba ngăn thuốc.",
    "Nếu các trường chỉ hiện '--' hoặc 'đang kết nối', ghi đúng là ảnh giao diện, không gọi đó là dữ liệu vận hành thực tế."
)
add_item(
    doc, "H05", "Hộp thoại cấu hình lịch và liều",
    "Trong mục 4.4.2, sau đoạn mô tả giao diện cập nhật cấu hình.",
    "Trên Dashboard bấm 'Cấu hình Lịch uống'. Chụp hộp thoại sao cho thấy đủ ngăn A/B/C, "
    "tên hoặc mã thuốc, giờ uống, số viên mỗi liều và trạng thái nạp. "
    "Không cần bấm Lưu chỉ để lấy ảnh.",
    "Hình: Giao diện cấu hình lịch uống và định lượng cho ba ngăn."
)
add_item(
    doc, "H06", "Cây dữ liệu Firebase Realtime Database",
    "Sau mục 3.4.2 hoặc trong mục 4.3.",
    "Mở Firebase Console của đúng dự án, vào Realtime Database → Data. "
    "Chụp các nhánh /config/compartments, /inventory/compartments, /deviceStatus và /events ở mức đủ đọc tên nhánh; "
    "có thể dùng hai ảnh nếu cây dữ liệu dài.",
    "Hình: Các nhánh dữ liệu cấu hình, tồn kho, trạng thái và sự kiện trên Firebase.",
    "Che hoặc loại khỏi khung hình token, khóa truy cập, email cá nhân và thông tin không cần cho báo cáo."
)

doc.add_heading("3. Ảnh minh chứng kiểm thử khi đã chạy ca tương ứng", level=2)
add_item(
    doc, "H07", "Phiên lấy thuốc đủ liều T05",
    "Trong mục 5.3 hoặc 5.4, cạnh dòng T05 của bảng kiểm thử.",
    "Ghi lại tồn kho và liều hiện tại của ngăn A trước thử. Trong Wokwi Serial Monitor gõ 1 để thử nhắc A, "
    "gõ T để TAKE, rồi gõ R đúng số viên theo cấu hình. Chụp LCD/LED/servo và log sau khi hệ thống phân loại kết quả; "
    "chờ giai đoạn đóng khoảng tám giây rồi chụp /events và /inventory trên Firebase để đối chiếu.",
    "Hình: Kết quả mô phỏng lấy đủ liều và bản ghi sự kiện, tồn kho sau phiên.",
    "Phím 1 là nhắc thử, không chứng minh lịch RTC/Firebase đã tự kích hoạt. Chỉ ghi PASS khi event và tồn kho thực tế khớp tiêu chí T05."
)
add_item(
    doc, "H08", "Các trường hợp thiếu, thừa liều và hoãn nhắc T06–T08",
    "Trong mục 5.3 hoặc 5.4, cạnh các dòng T06–T08 đã thực hiện.",
    "Chạy từng ca riêng sau khi trở về IDLE: T06 dùng TAKE rồi REMOVE ít hơn liều hoặc chờ hết 60 giây; "
    "T07 dùng TAKE rồi REMOVE nhiều hơn liều; T08 dùng SNOOZE và theo dõi thời điểm nhắc lại/bỏ lỡ. "
    "Mỗi ca chụp trạng thái Wokwi, log và event Firebase tương ứng, kèm thời điểm bắt đầu/kết thúc.",
    "Hình: Cảnh báo và event của ca kiểm thử [ghi đúng mã T06, T07 hoặc T08].",
    "Không gộp ảnh của các lần chạy khác nhau thành một kết quả PASS; T08 cần quan sát đủ khoảng thời gian nêu trong ca thử."
)
add_item(
    doc, "H09", "Lệnh chỉnh tồn kho T11",
    "Trong mục 4.3.1 và dòng T11 của mục 5.3/5.4.",
    "Ghi giá trị tồn kho trước thử. Trên Dashboard mở cấu hình ngăn, nhập số tồn kho đích và lưu. "
    "Chụp trạng thái command đang chờ; sau khi ESP32 xử lý, chụp status applied hoặc rejected, "
    "/inventory/compartments trên Firebase và số mới trên Dashboard.",
    "Hình: Lệnh điều chỉnh tồn kho từ Dashboard và kết quả ESP32 xác nhận.",
    "Firebase nhận yêu cầu chưa đồng nghĩa ESP32 đã áp dụng; phải có trạng thái command và tồn kho thiết bị công bố."
)
add_item(
    doc, "H10", "Nguồn giờ và hoạt động khi mất mạng T01–T04, T12",
    "Trong mục 5.3/5.4, chỉ cạnh ca đã chạy.",
    "Với T01, chụp LCD/Serial thể hiện NTP hoặc RTC và giờ ICT; ảnh 04 đã có cho thấy log NTP và tải cấu hình. "
    "Với T02–T04 hoặc T12, chuẩn bị điều kiện ban đầu đúng bảng thử, ngắt mạng có kiểm soát, "
    "chụp log trước/sau cùng dữ liệu Firebase khi kết nối lại.",
    "Hình: Trạng thái thời gian hoặc đồng bộ ngoại tuyến của ca kiểm thử [ghi mã Txx].",
    "Một ảnh log NTP không đủ kết luận toàn bộ các ca offline đã đạt."
)

doc.add_heading("4. Minh chứng chỉ thu thập khi có phần cứng thật", level=2)
doc.add_paragraph(
    "T14 cần ảnh board ESP32, cách nối GPIO35 và nguồn servo kèm log/đo điện áp khi thao tác; "
    "T15 cần ảnh mạch UPS/power-path, dụng cụ đo, biểu đồ hoặc bảng điện áp–dòng–thời gian trong ít nhất bốn giờ. "
    "Sơ đồ Wokwi không thay thế được hai phép thử này. Nếu nhóm chưa lắp và đo, giữ trạng thái 'chưa kiểm thử', không chèn ảnh minh họa như kết quả."
)

doc.add_heading("5. Cách lưu ảnh để dễ đối chiếu", level=2)
doc.add_paragraph(
    "Mỗi ảnh nên có tên gồm mã hình hoặc mã test, ngày giờ và nội dung, ví dụ H07_T05_2026-10-18_event.png. "
    "Cùng ảnh lưu một dòng ghi commit, cấu hình liều/tồn kho ban đầu, thao tác, kết quả quan sát và PASS/FAIL. "
    "Khi dán vào Word, dùng Insert → Pictures → This Device, đặt chú thích ngay dưới ảnh và giữ ảnh đủ lớn để đọc LCD/log. "
    "Chỉ cắt phần màn hình không liên quan; không chỉnh sửa số liệu hay che mất thông báo lỗi."
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
